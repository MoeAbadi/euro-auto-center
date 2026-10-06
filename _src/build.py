#!/usr/bin/env python3
"""Builds the Euro Auto Center static site into ../site_out."""
import html, json, os, shutil, sys
from urllib.parse import quote_plus

sys.path.insert(0, os.path.dirname(__file__))
from brands import BRANDS, SYSTEMS

SRC = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(os.path.dirname(SRC), "site_out")
SITE = "https://www.euroautocenterllc.com"
PDF = "euro-auto-center_profile.pdf"
TEMPLATE_XLSX = "assets/files/quotation-request-template.xlsx"  # maintained by hand in the repo; not generated
E = html.escape
import hashlib
def _v(name):
    return hashlib.md5(open(os.path.join(SRC, name), "rb").read()).hexdigest()[:8]
ASSET_V = {"css": _v("site.css"), "js": _v("site.js")}

SALES = "sales@euro-auto-center.com"
DIRECT = "w.saad@euro-auto-center.com"

OFFICES = [
    dict(key="rak", label="Headquarters", city="Ras Al Khaimah, UAE",
         lines=["Al Shohada Road, Compass Building FDRK2383", "Ras Al Khaimah – UAE"],
         map_q="Compass Building, Al Shohada Road, Ras Al Khaimah, United Arab Emirates",
         ld=dict(street="Al Shohada Road, Compass Building FDRK2383", city="Ras Al Khaimah", country="AE")),
]

# ---------- icons (generic UI icons) ----------
def ic(path, sw="2"):
    return f'<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="{sw}" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true">{path}</svg>'

I = dict(
    download=ic('<path d="M12 3v12M7 10l5 5 5-5M5 21h14"/>'),
    arrow=ic('<path d="M5 12h14M13 6l6 6-6 6"/>'),
    ext=ic('<path d="M14 4h6v6M20 4l-9 9M18 14v5a1 1 0 0 1-1 1H5a1 1 0 0 1-1-1V7a1 1 0 0 1 1-1h5"/>'),
    check=ic('<path d="M5 12l5 5L20 7"/>', "3"),
    shield=ic('<path d="M12 2l8 4v6c0 5-3.5 8.5-8 10-4.5-1.5-8-5-8-10V6l8-4z"/><path d="M9 12l2 2 4-4"/>', "1.8"),
    target=ic('<circle cx="12" cy="12" r="9"/><circle cx="12" cy="12" r="5"/><circle cx="12" cy="12" r="1"/>', "1.8"),
    pin=ic('<path d="M21 10c0 7-9 13-9 13S3 17 3 10a9 9 0 0 1 18 0z"/><circle cx="12" cy="10" r="3"/>'),
    globe=ic('<circle cx="12" cy="12" r="9"/><path d="M3 12h18M12 3a14 14 0 0 1 0 18M12 3a14 14 0 0 0 0 18"/>', "1.8"),
    search=ic('<circle cx="11" cy="11" r="7"/><path d="M21 21l-4.3-4.3"/>'),
    db=ic('<ellipse cx="12" cy="5" rx="8" ry="3"/><path d="M4 5v6c0 1.66 3.58 3 8 3s8-1.34 8-3V5"/><path d="M4 11v6c0 1.66 3.58 3 8 3s8-1.34 8-3v-6"/>', "1.8"),
    truck=ic('<path d="M1 6h13v11H1zM14 10h4l3 3v4h-7z"/><circle cx="5.5" cy="18.5" r="1.8"/><circle cx="17.5" cy="18.5" r="1.8"/>', "1.8"),
    gear=ic('<circle cx="12" cy="12" r="3"/><path d="M12 2v3M12 19v3M4.93 4.93l2.12 2.12M16.95 16.95l2.12 2.12M2 12h3M19 12h3M4.93 19.07l2.12-2.12M16.95 7.05l2.12-2.12"/>', "1.8"),
    chart=ic('<path d="M3 3v18h18"/><path d="M7 15l4-4 3 3 5-6"/>', "1.8"),
    mail=ic('<rect x="2" y="4" width="20" height="16" rx="2"/><path d="M22 7l-10 6L2 7"/>'),
    user=ic('<circle cx="12" cy="8" r="4"/><path d="M4 21c0-4 3.6-7 8-7s8 3 8 7"/>'),
    phone=ic('<path d="M22 16.92v3a2 2 0 0 1-2.18 2 19.8 19.8 0 0 1-8.63-3.07 19.5 19.5 0 0 1-6-6 19.8 19.8 0 0 1-3.07-8.67A2 2 0 0 1 4.11 2h3a2 2 0 0 1 2 1.72c.13.96.36 1.9.7 2.81a2 2 0 0 1-.45 2.11L8.09 9.91a16 16 0 0 0 6 6l1.27-1.27a2 2 0 0 1 2.11-.45c.91.34 1.85.57 2.81.7A2 2 0 0 1 22 16.92z"/>'),
    chat=ic('<path d="M21 12a8 8 0 0 1-11.6 7.1L4 20l1-4.6A8 8 0 1 1 21 12z"/>'),
    file=ic('<path d="M14 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V8z"/><path d="M14 2v6h6"/>'),
    sheet=ic('<path d="M14 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V8z"/><path d="M14 2v6h6M8 13h8M8 17h8M12 11v8"/>'),
    book=ic('<path d="M4 19.5A2.5 2.5 0 0 1 6.5 17H20V3H6.5A2.5 2.5 0 0 0 4 5.5z"/><path d="M4 19.5A2.5 2.5 0 0 0 6.5 22H20v-5"/>'),
    upload=ic('<path d="M12 16V4M7 9l5-5 5 5M4 20h16"/>', "1.8"),
    trash=ic('<path d="M3 6h18M8 6V4h8v2M6 6l1 14h10l1-14"/>'),
    plus=ic('<path d="M12 5v14M5 12h14"/>'),
    route=ic('<circle cx="6" cy="19" r="2"/><circle cx="18" cy="5" r="2"/><path d="M8 19h8a4 4 0 0 0 0-8H8a4 4 0 0 1 0-8h8"/>'),
    users=ic('<circle cx="9" cy="8" r="3.5"/><path d="M2 21c0-3.5 3-6 7-6s7 2.5 7 6M16 4.5a3.5 3.5 0 0 1 0 7M22 21c0-3-2-5.2-5-5.8"/>'),
    camera=ic('<rect x="3" y="3" width="18" height="18" rx="5"/><circle cx="12" cy="12" r="4"/>'),
    briefcase=ic('<rect x="2" y="7" width="20" height="14" rx="2"/><path d="M16 7V5a2 2 0 0 0-2-2h-4a2 2 0 0 0-2 2v2"/>'),
)
MENU_OPEN = '<svg class="icon-open" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" aria-hidden="true"><path d="M4 7h16M4 12h16M4 17h16"/></svg>'
MENU_CLOSE = '<svg class="icon-close" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" hidden aria-hidden="true"><path d="M6 6l12 12M18 6L6 18"/></svg>'
GEAR_BIG = '<svg class="hero__gear" viewBox="0 0 200 200" aria-hidden="true" fill="currentColor"><path d="M113 8l4 22a72 72 0 0 1 18 8l19-13 18 18-13 19a72 72 0 0 1 8 18l22 4v26l-22 4a72 72 0 0 1-8 18l13 19-18 18-19-13a72 72 0 0 1-18 8l-4 22H87l-4-22a72 72 0 0 1-18-8l-19 13-18-18 13-19a72 72 0 0 1-8-18L11 113V87l22-4a72 72 0 0 1 8-18L28 46l18-18 19 13a72 72 0 0 1 18-8l4-22h26zM100 62a38 38 0 1 0 0 76 38 38 0 0 0 0-76z"/></svg>'

SOCIALS = [("linkedin", "LinkedIn", "briefcase"), ("facebook", "Facebook", "users"), ("instagram", "Instagram", "camera")]

NAV = [("", "Home"), ("about-us/", "About Us"), ("brands-catalogs/", "Brands &amp; Catalogs"), ("contact/", "Contact")]

# ---------- helpers ----------
def tier_cls(t):
    return {"Premium Aftermarket": "tier tier--premium", "Value Line": "tier tier--value"}.get(t, "tier")

def logo_file(slug):
    slug = slug.lower()
    for ext in ("svg", "png", "webp"):
        if os.path.exists(os.path.join(SRC, "logos", f"{slug}.{ext}")):
            return f"assets/brands/{slug}.{ext}"
    return None

def wordmark(b, r, tag="div", href=None, cls="wordmark"):
    lf = logo_file(b.get("logo", b["slug"]))
    inner = f'<img src="{r}{lf}" alt="{E(b["name"])} logo" loading="lazy">' if lf else f'<span>{E(b["name"])}</span>'
    if href:
        return f'<a class="{cls}" href="{href}" aria-label="{E(b["name"])}">{inner}</a>'
    return f'<{tag} class="{cls}" aria-hidden="true">{inner}</{tag}>'

def social_links(r):
    out = []
    for key, label, icon in SOCIALS:
        f = next((f"assets/social/{key}.{x}" for x in ("svg", "png") if os.path.exists(os.path.join(SRC, "social", f"{key}.{x}"))), None)
        icon_html = f'<img src="{r}{f}" alt="">' if f else I[icon]
        out.append(f'<a data-social="{key}" aria-disabled="true" title="{label} page coming soon">{icon_html}{label}</a>')
    return '<div class="social">' + "".join(out) + "</div>"

def ld_org():
    locs = []
    for o in OFFICES:
        locs.append({
            "@type": "AutoPartsStore", "name": f"Euro Auto Center – {o['city']}",
            "address": {"@type": "PostalAddress", "streetAddress": o["ld"]["street"],
                        "addressLocality": o["ld"]["city"], "addressCountry": o["ld"]["country"]},
            "email": SALES})
    data = {"@context": "https://schema.org", "@type": "Organization", "name": "Euro Auto Center",
            "url": SITE + "/", "logo": SITE + "/eac-logo.png", "email": SALES,
            "areaServed": ["Middle East", "Africa"], "department": locs}
    return '<script type="application/ld+json">' + json.dumps(data, ensure_ascii=False) + "</script>"

def page(path, title, desc, body, active="", extra_head=""):
    depth = path.count("/")
    r = "../" * depth
    canon = SITE + "/" + (path[:-10] if path.endswith("index.html") else path)
    cur = ' aria-current="page"'
    nav_items = "".join(
        f'<li><a href="{r}{href}"{cur if active == href else ""}>{label}</a></li>'
        for href, label in NAV)
    head = f'''<!DOCTYPE html>
<html lang="en" class="no-js">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<meta name="description" content="{E(desc)}">
<link rel="canonical" href="{canon}">
<meta name="theme-color" content="#0a1628">
<meta property="og:title" content="{title}">
<meta property="og:description" content="{E(desc)}">
<meta property="og:type" content="website">
<meta property="og:url" content="{canon}">
<meta property="og:image" content="{SITE}/eac-logo.png">
<link rel="icon" type="image/png" href="{r}eac-logo.png">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600&family=Sora:wght@600;700;800&display=swap" rel="stylesheet">
<link rel="stylesheet" href="{r}assets/css/site.css?v={ASSET_V['css']}">
{extra_head}</head>
<body>
<a class="skip" href="#main">Skip to content</a>
<header class="header" id="header">
  <div class="container header__inner">
    <a href="{r or './'}" class="logo" aria-label="Euro Auto Center home">
      <span class="logo__mark"><img src="{r}eac-logo.png" alt="EAC" width="65" height="24"></span>
      <span class="logo__text">EURO <span>AUTO</span><br>CENTER</span>
    </a>
    <nav class="nav" id="nav" aria-label="Main">
      <ul class="nav__links">{nav_items}</ul>
      <a href="{r}request-quote/" class="btn btn--primary">Request a Quote</a>
    </nav>
    <button class="nav-toggle" id="navToggle" aria-label="Open menu" aria-expanded="false" aria-controls="nav">{MENU_OPEN}{MENU_CLOSE}</button>
  </div>
</header>
<main id="main">
'''
    foot = f'''</main>
<footer class="footer">
  <div class="container">
    <div class="footer__grid">
      <div>
        <a href="{r or './'}" class="logo" aria-label="Euro Auto Center home">
          <span class="logo__mark"><img src="{r}eac-logo.png" alt="EAC" width="65" height="24"></span>
          <span class="logo__text">EURO <span>AUTO</span><br>CENTER</span>
        </a>
        <p class="footer__tagline">Genuine European OEM and premium aftermarket parts, backed by accurate parts data, for the Middle East &amp; Africa.</p>
        {social_links(r)}
      </div>
      <div>
        <h4>Company</h4>
        <ul>
          <li><a href="{r}about-us/">About Us</a></li>
          <li><a href="{r}brands-catalogs/">Brands &amp; Catalogs</a></li>
          <li><a href="{r}request-quote/">Request a Quote</a></li>
          <li><a href="{r}downloads/">Downloads</a></li>
          <li><a href="{r}contact/">Contact</a></li>
        </ul>
      </div>
      <div>
        <h4>Contact</h4>
        <ul>
          <li><a href="mailto:{SALES}">{SALES}</a></li>
          <li><a href="mailto:{DIRECT}">{DIRECT}</a></li>
        </ul>
      </div>
      <div>
        <h4>Headquarters</h4>
        <ul>
          {"".join(f'<li><b style="color:#fff">{o["label"]}</b><br>{"<br>".join(E(l) for l in o["lines"])}</li>' for o in OFFICES)}
        </ul>
      </div>
    </div>
    <div class="footer__bottom">
      <span>© <span id="year">2026</span> Euro Auto Center. All rights reserved.</span>
      <span>Brand names and logos are trademarks of their respective owners.</span>
    </div>
  </div>
</footer>
<script src="{r}assets/js/site.js?v={ASSET_V['js']}" defer></script>
</body>
</html>
'''
    full = head + body.replace("{R}", r) + foot
    dest = os.path.join(OUT, path)
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    with open(dest, "w", encoding="utf-8") as f:
        f.write(full)
    return canon

def cta_profile():
    return f'''
<section class="section section--tight">
  <div class="container">
    <div class="cta reveal">
      <div>
        <span class="eyebrow">Company Profile</span>
        <h2>Get the full picture of Euro Auto Center.</h2>
        <p>Download our company profile for our specialization, service territory, and how to work with us.</p>
      </div>
      <div class="cta__side">
        <a href="{{R}}{PDF}" class="btn btn--primary btn--lg" download>{I["download"]}Download Company Profile</a>
        <span class="cta__meta">PDF · 100 KB</span>
      </div>
    </div>
  </div>
</section>'''

def brand_card(b, r):
    tags = "".join(f'<span class="tag">{SYSTEMS[s]}</span>' for s in b["systems"])
    search = " ".join([b["name"], b["group"], b["hq"], b["tier"]] + [SYSTEMS[s] for s in b["systems"]]).lower()
    nm = E(b["name"])
    cat = (f'<a class="btn btn--outline" href="{b["catalog"]}" target="_blank" rel="noopener" aria-label="View {nm} catalog (opens in new tab)">{I["ext"]}Catalog</a>'
           if b["catalog"] else f'<a class="btn btn--outline" href="{r}contact/" aria-label="{nm} catalog on request">{I["book"]}On request</a>')
    group = f'<span><b>Group</b> {E(b["group"])}</span>' if b["group"] else ""
    return f'''<article class="brand-card" data-systems="{" ".join(b["systems"])}" data-search="{E(search)}">
  {wordmark(b, r, href=f"{r}brands/{b['slug']}/")}
  <div class="brand-card__body">
    <span class="{tier_cls(b["tier"])}">{b["tier"]}</span>
    <h3><a href="{r}brands/{b['slug']}/">{E(b["name"])}</a></h3>
    <div class="meta">{group}<span><b>HQ</b> {E(b["hq"])}</span></div>
    <div class="tags">{tags}</div>
    <div class="brand-card__actions">
      {cat}
      <a class="btn btn--primary" href="{r}request-quote/?brand={b['slug']}">Inquire</a>
    </div>
  </div>
</article>'''

# ---------- pages ----------
def build():
    if os.path.exists(OUT):
        shutil.rmtree(OUT)
    os.makedirs(OUT)
    urls = []

    # ---- Home ----
    marquee_items = "".join(wordmark(b, "{R}", href=f"{{R}}brands/{b['slug']}/") for b in BRANDS)
    home = f'''
<section class="hero">
  <div class="hero__glow" aria-hidden="true"></div>
  {GEAR_BIG}
  <div class="container">
    <div class="hero__content">
      <span class="eyebrow">European Auto Parts · Middle East &amp; Africa</span>
      <h1>Your trusted <em>European</em> parts partner.</h1>
      <p class="hero__lead">Genuine OEM and premium aftermarket parts from {len(BRANDS)} leading European and global brands, identified accurately and supplied across the Middle East &amp; Africa.</p>
      <div class="hero__actions">
        <a href="{{R}}request-quote/" class="btn btn--primary btn--lg">Request a Quote {I["arrow"]}</a>
        <a href="{{R}}{PDF}" class="btn btn--ghost btn--lg" download>{I["download"]}Company Profile</a>
      </div>
    </div>
    <div class="badges">
      <div class="badge"><span class="badge__icon">{I["shield"]}</span><div><strong>Genuine Parts</strong><span>Authorized brands, full traceability</span></div></div>
      <div class="badge"><span class="badge__icon">{I["target"]}</span><div><strong>OE-Matched</strong><span>Identified by OE number and VIN</span></div></div>
      <div class="badge"><span class="badge__icon">{I["pin"]}</span><div><strong>UAE Based</strong><span>Headquartered in Ras Al Khaimah, UAE</span></div></div>
      <div class="badge"><span class="badge__icon">{I["globe"]}</span><div><strong>Regional Reach</strong><span>Supplying the Middle East &amp; Africa</span></div></div>
    </div>
  </div>
</section>

<section class="section">
  <div class="container">
    <div class="section-head section-head--row reveal">
      <div>
        <span class="eyebrow">Authorized Brands</span>
        <h2>{len(BRANDS)} brands you already trust.</h2>
        <p>From braking and filtration to clutches, suspension and engine timing, sourced from the manufacturers that supply the car makers.</p>
      </div>
      <a href="{{R}}brands-catalogs/" class="btn btn--outline">Browse all brands {I["arrow"]}</a>
    </div>
  </div>
  <div class="marquee" aria-label="Authorized brands">
    <div class="marquee__track">{marquee_items}{marquee_items.replace('<a ', '<a tabindex="-1" aria-hidden="true" ')}</div>
  </div>
</section>

<section class="section section--soft">
  <div class="container">
    <div class="section-head reveal">
      <span class="eyebrow">What We Do</span>
      <h2>Parts supply and parts data, under one roof.</h2>
      <p>We cover the full parts cycle for European vehicles, from a single hard-to-find component to clean, ERP-ready catalog data.</p>
    </div>
    <div class="grid-3">
      <article class="card card--hover reveal"><div class="card__icon">{I["gear"]}</div><h3>European OEM &amp; Aftermarket</h3><p>Genuine original-equipment and approved premium aftermarket parts from authorized brands.</p><a class="link" href="{{R}}brands-catalogs/">See brands {I["arrow"]}</a></article>
      <article class="card card--hover reveal"><div class="card__icon">{I["search"]}</div><h3>Data-Driven Identification</h3><p>OE-number and VIN-based identification with cross-referencing, so you get the right part the first time.</p><a class="link" href="{{R}}request-quote/">Send part numbers {I["arrow"]}</a></article>
      <article class="card card--hover reveal"><div class="card__icon">{I["truck"]}</div><h3>Regional Distribution</h3><p>B2B supply for workshops, dealers and distributors across the Middle East &amp; Africa.</p><a class="link" href="{{R}}contact/">Talk to sales {I["arrow"]}</a></article>
    </div>
  </div>
</section>

<section class="section">
  <div class="container">
    <div class="section-head reveal">
      <span class="eyebrow">Systems We Cover</span>
      <h2>Find brands by product system.</h2>
    </div>
    <div class="chips reveal">{"".join(f'<a class="chip" href="{{R}}brands-catalogs/?system={k}">{v}</a>' for k, v in SYSTEMS.items())}</div>
  </div>
</section>
{cta_profile()}
'''
    urls.append(page("index.html", "Euro Auto Center | European OEM & Aftermarket Auto Parts",
                     f"Euro Auto Center supplies genuine European OEM and premium aftermarket parts from {len(BRANDS)} authorized brands across the Middle East & Africa from Ras Al Khaimah, UAE.",
                     home, "", ld_org()))

    # ---- About ----
    offices_html = "".join(f'''<div class="office"><span class="label">{o["label"]}</span><h3>{o["city"]}</h3><p>{"<br>".join(E(l) for l in o["lines"])}</p></div>''' for o in OFFICES)
    about = f'''
<section class="page-hero">
  <div class="container">
    <nav class="crumbs" aria-label="Breadcrumb"><a href="{{R}}">Home</a><span>/</span><span>About Us</span></nav>
    <h1>About Euro Auto Center</h1>
    <p>A European auto parts distributor serving workshops, dealers and distributors across the Middle East &amp; Africa.</p>
  </div>
</section>
<section class="section">
  <div class="container about">
    <div class="about__text reveal">
      <span class="eyebrow">Who We Are</span>
      <h2 style="font-size:clamp(28px,4vw,38px);margin:14px 0 20px">Genuine European parts, delivered with precision.</h2>
      <p>Euro Auto Center is headquartered in Ras Al Khaimah, United Arab Emirates. We specialize in European OEM and quality aftermarket parts and supply customers across the Middle East &amp; Africa.</p>
      <p>We combine a strong supply network of {len(BRANDS)} authorized brands with structured parts data, so every order is identified correctly, sourced quickly, and delivered with full traceability.</p>
      <ul class="checklist">
        <li><span class="check">{I["check"]}</span>Genuine and quality-approved brands only</li>
        <li><span class="check">{I["check"]}</span>Accurate part identification by OE number and VIN</li>
        <li><span class="check">{I["check"]}</span>UAE headquarters serving the Gulf, Middle East &amp; Africa</li>
      </ul>
    </div>
    <div class="reveal">{offices_html}</div>
  </div>
</section>
<section class="section section--soft">
  <div class="container">
    <div class="section-head reveal"><span class="eyebrow">Our Services</span><h2>Everything around the part.</h2></div>
    <div class="grid-3">
      <article class="card reveal"><div class="card__icon">{I["gear"]}</div><h3>OEM &amp; Aftermarket Supply</h3><p>Original-equipment and premium aftermarket parts for European, Asian and American vehicles.</p></article>
      <article class="card reveal"><div class="card__icon">{I["search"]}</div><h3>Parts Sourcing</h3><p>Fast location of hard-to-find and specialist components, using OE-number and VIN identification.</p></article>
      <article class="card reveal"><div class="card__icon">{I["db"]}</div><h3>Parts Data &amp; Cataloging</h3><p>Cross-referencing, catalog cleansing and ERP-ready item records for your parts master data.</p></article>
      <article class="card reveal"><div class="card__icon">{I["chart"]}</div><h3>Inventory &amp; Reporting</h3><p>Stock analysis and demand insights that help partners plan purchasing and reduce dead stock.</p></article>
      <article class="card reveal"><div class="card__icon">{I["truck"]}</div><h3>Regional Distribution</h3><p>Reliable B2B supply and dispatch to workshops, dealers and distributors.</p></article>
      <article class="card reveal"><div class="card__icon">{I["users"]}</div><h3>Dedicated Account Contact</h3><p>One point of contact for quotes, follow-up and after-sales questions.</p></article>
    </div>
  </div>
</section>
<section class="section">
  <div class="container">
    <div class="section-head reveal"><span class="eyebrow">Why Euro Auto Center</span><h2>Built on precision, reliability and speed.</h2></div>
    <div class="why reveal">
      <div class="why__item"><div class="why__num">01</div><h3>Precision</h3><p>The right part the first time, through data-driven identification and cross-referencing.</p></div>
      <div class="why__item"><div class="why__num">02</div><h3>Reliability</h3><p>Genuine and quality-approved brands only, with consistent supply you can plan around.</p></div>
      <div class="why__item"><div class="why__num">03</div><h3>Speed</h3><p>Quick quotations and dispatch to keep your vehicles on the road.</p></div>
      <div class="why__item"><div class="why__num">04</div><h3>Partnership</h3><p>Long-term relationships with workshops and distributors across the region.</p></div>
    </div>
  </div>
</section>
{cta_profile()}
'''
    urls.append(page("about-us/index.html", "About Us | Euro Auto Center",
                     "Euro Auto Center is a European auto parts distributor headquartered in Ras Al Khaimah, United Arab Emirates.",
                     about, "about-us/"))

    # ---- Brands hub ----
    cards = "".join(brand_card(b, "{R}") for b in sorted(BRANDS, key=lambda b: b["name"].lower()))
    chips = '<button type="button" class="chip" data-system="all" aria-pressed="true">All systems</button>' + "".join(
        f'<button type="button" class="chip" data-system="{k}" aria-pressed="false">{v}</button>' for k, v in SYSTEMS.items())
    hub = f'''
<section class="page-hero">
  <div class="container">
    <nav class="crumbs" aria-label="Breadcrumb"><a href="{{R}}">Home</a><span>/</span><span>Brands &amp; Catalogs</span></nav>
    <h1>Brands &amp; Catalogs</h1>
    <p>Browse our {len(BRANDS)} authorized brands, open each manufacturer's official online catalog, or send us an inquiry.</p>
  </div>
</section>
<section class="section">
  <div class="container">
    <div class="toolbar">
      <label class="search" for="brandSearch">{I["search"]}
        <input id="brandSearch" type="search" placeholder="Search by brand, group or system (e.g. clutch, MANN, ZF)" autocomplete="off" aria-label="Search brands">
      </label>
      <div class="chips" role="group" aria-label="Filter by product system">{chips}</div>
      <p class="result-count" id="resultCount" aria-live="polite"></p>
    </div>
    <div class="brand-grid" id="brandGrid">{cards}</div>
    <div class="empty" id="emptyState" hidden>
      <p>No brands match your search.</p>
      <p style="margin-top:12px"><button type="button" class="btn btn--outline btn--sm" id="resetFilters">Clear filters</button></p>
    </div>
    <p class="note">"View Catalog" opens the manufacturer's official online catalog in a new tab. Brand names and logos are trademarks of their respective owners. Can't find a part? <a href="{{R}}request-quote/" style="color:var(--blue-500);font-weight:600">Send us the OE number</a>.</p>
  </div>
</section>
'''
    urls.append(page("brands-catalogs/index.html", "Brands & Catalogs | Euro Auto Center",
                     f"Browse {len(BRANDS)} authorized brands including Bosch, MANN-FILTER, MAHLE, febi, LuK, SACHS, TRW and Valeo. Search by product system and open official online catalogs.",
                     hub, "brands-catalogs/"))

    # ---- Brand detail pages ----
    for b in BRANDS:
        r = "../../"
        related = [x for x in BRANDS if x is not b and x["group"] and x["group"] == b["group"]]
        if len(related) < 3:
            related += [x for x in BRANDS if x is not b and x not in related and set(x["systems"]) & set(b["systems"])]
        related = related[:4]
        rel_html = "".join(brand_card(x, "{R}") for x in related)
        sys_tags = "".join(f'<a class="chip" href="{{R}}brands-catalogs/?system={s}">{SYSTEMS[s]}</a>' for s in b["systems"])
        cat_btn = (f'<a class="btn btn--ghost btn--lg" href="{b["catalog"]}" target="_blank" rel="noopener">{I["ext"]}View {E(b["name"])} Catalog</a>'
                   if b["catalog"] else f'<a class="btn btn--ghost btn--lg" href="{{R}}contact/">{I["book"]}Catalog on request</a>')
        group_fact = f'<div><dt>Group</dt><dd>{E(b["group"])}</dd></div>' if b["group"] else '<div><dt>Group</dt><dd>—</dd></div>'
        body = f'''
<section class="page-hero">
  <div class="container">
    <nav class="crumbs" aria-label="Breadcrumb"><a href="{{R}}">Home</a><span>/</span><a href="{{R}}brands-catalogs/">Brands &amp; Catalogs</a><span>/</span><span>{E(b["name"])}</span></nav>
    <div class="brand-hero" style="margin-top:22px">
      {wordmark(b, "{R}")}
      <div>
        <span class="{tier_cls(b["tier"])}">{b["tier"]}</span>
        <h1>{E(b["name"])}</h1>
        <p>{E(b["desc"])}</p>
        <div class="hero__actions" style="margin-top:26px">
          <a class="btn btn--primary btn--lg" href="{{R}}request-quote/?brand={b['slug']}">Inquire about {E(b["name"])} {I["arrow"]}</a>
          {cat_btn}
        </div>
      </div>
    </div>
  </div>
</section>
<section class="section section--tight">
  <div class="container">
    <dl class="facts reveal">
      {group_fact}
      <div><dt>Group headquarters</dt><dd>{E(b["hq"])}</dd></div>
      <div><dt>Quality tier</dt><dd>{b["tier"]}</dd></div>
    </dl>
    <h2 style="font-size:22px;margin:40px 0 14px">Product systems</h2>
    <div class="chips">{sys_tags}</div>
  </div>
</section>
<section class="section section--soft section--tight">
  <div class="container">
    <div class="section-head reveal"><span class="eyebrow">How to order</span><h2>Get {E(b["name"])} parts in three steps.</h2></div>
    <ol class="steps reveal">
      <li><h3>Find the part</h3><p>Look it up in the {E(b["name"])} online catalog, or simply send us the OE number or your VIN.</p></li>
      <li><h3>Request a quote</h3><p>Send the part numbers and quantities through our quote form or by email.</p></li>
      <li><h3>Confirm &amp; receive</h3><p>We confirm availability, price and lead time, then dispatch to you.</p></li>
    </ol>
  </div>
</section>
<section class="section">
  <div class="container">
    <div class="section-head section-head--row reveal"><div><span class="eyebrow">Related</span><h2>Related brands</h2></div><a class="btn btn--outline" href="{{R}}brands-catalogs/">All brands {I["arrow"]}</a></div>
    <div class="brand-grid">{rel_html}</div>
  </div>
</section>
'''
        urls.append(page(f"brands/{b['slug']}/index.html", f"{b['name']} Parts | Euro Auto Center",
                         f"{b['name']} {b['tier'].lower()} parts from Euro Auto Center: {', '.join(SYSTEMS[s] for s in b['systems'])}. Request a quote or open the official {b['name']} catalog.",
                         body, "brands-catalogs/"))

    # ---- Request a Quote ----
    brand_opts = '<option value="">Any / best offer</option>' + "".join(
        f'<option value="{b["slug"]}">{E(b["name"])}</option>' for b in sorted(BRANDS, key=lambda b: b["name"].lower()))
    line_tpl = f'''<div class="line">
  <div class="field"><label data-n="lOe">OE number</label><input data-n="lOe" type="text" autocomplete="off" placeholder="e.g. 1K0 615 301 AA"></div>
  <div class="field"><label data-n="lPn">Part number</label><input data-n="lPn" type="text" autocomplete="off" placeholder="Brand part no."></div>
  <div class="field"><label data-n="lBr">Preferred brand</label><select data-n="lBr">{brand_opts}</select></div>
  <div class="field field--qty"><label data-n="lQty">Qty</label><input data-n="lQty" type="number" min="1" step="1" value="1" inputmode="numeric"></div>
  <button type="button" class="icon-btn" aria-label="Remove line">{I["trash"]}</button>
</div>'''
    quote = f'''
<section class="page-hero">
  <div class="container">
    <nav class="crumbs" aria-label="Breadcrumb"><a href="{{R}}">Home</a><span>/</span><span>Request a Quote</span></nav>
    <h1>Request a Quote</h1>
    <p>Send us OE or part numbers, quantities and your VIN. We reply with availability, price and lead time.</p>
  </div>
</section>
<section class="section">
  <div class="container form-wrap">
    <form class="panel" id="quoteForm" novalidate>
      <fieldset>
        <legend>Your details</legend>
        <div class="fields">
          <div class="field"><label for="qName">Full name *</label><input id="qName" type="text" autocomplete="name" required></div>
          <div class="field"><label for="qCompany">Company <small>(optional)</small></label><input id="qCompany" type="text" autocomplete="organization"></div>
          <div class="field"><label for="qEmail">Email *</label><input id="qEmail" type="email" autocomplete="email" required></div>
          <div class="field"><label for="qCountry">Country</label><input id="qCountry" type="text" autocomplete="country-name"></div>
        </div>
      </fieldset>
      <fieldset>
        <legend>Vehicle <small class="muted" style="font-weight:400;font-family:var(--font-body)">(optional, helps us match the right part)</small></legend>
        <div class="fields">
          <div class="field field--full"><label for="qVin">VIN (chassis number)</label><input id="qVin" type="text" maxlength="17" autocomplete="off" style="text-transform:uppercase" placeholder="17 characters"></div>
          <div class="field"><label for="qMake">Make</label><input id="qMake" type="text" placeholder="e.g. Volkswagen"></div>
          <div class="field"><label for="qModel">Model &amp; year</label><input id="qModel" type="text" placeholder="e.g. Golf 7, 2016"></div>
          <input id="qYear" type="hidden" value="">
        </div>
      </fieldset>
      <fieldset>
        <legend>Parts</legend>
        <div class="lines" id="lines"></div>
        <button type="button" class="btn btn--outline btn--sm add-line" id="addLine">{I["plus"]}Add another part</button>
      </fieldset>
      <fieldset>
        <legend>Bulk list &amp; notes</legend>
        <label class="check-row" for="qAttach"><input id="qAttach" type="checkbox"><span>I will attach an Excel / CSV parts list to the email</span></label>
        <div class="field" style="margin-top:16px"><label for="qNotes">Notes</label><textarea id="qNotes" placeholder="Delivery destination, required dates, or anything else we should know"></textarea></div>
      </fieldset>
      <div class="form-actions">
        <button type="submit" class="btn btn--primary btn--lg">{I["mail"]}Send by Email</button>
        <button type="button" class="btn btn--outline" id="copyBtn" hidden>Copy request</button>
      </div>
      <p class="form-msg" id="formMsg" role="status" aria-live="polite"></p>
      <p class="muted" style="font-size:13px;margin-top:6px">"Send by Email" opens your email app with the request filled in, addressed to {SALES}. Nothing is sent until you press Send.</p>
    </form>
    <aside>
      <div class="bulk">
        {I["upload"]}
        <h3>Have a long parts list?</h3>
        <p>Download our Excel template, fill in the OE numbers and quantities, tick "I will attach…" and attach it to the email.</p>
        <a class="btn btn--primary btn--sm" href="{{R}}{TEMPLATE_XLSX}" download>{I["sheet"]}Download Excel template</a>
        <p style="margin:12px 0 0;font-size:13px">Excel (.xlsx, .xls) and CSV files accepted.</p>
      </div>
      <div class="panel" style="margin-top:20px">
        <h2 style="font-size:18px;margin-bottom:14px">Prefer to talk?</h2>
        <ul class="side-list">
          <li>{I["mail"]}<a href="mailto:{SALES}">{SALES}</a></li>
        </ul>
      </div>
    </aside>
  </div>
  <template id="lineTpl">{line_tpl}</template>
</section>
'''
    urls.append(page("request-quote/index.html", "Request a Quote | Euro Auto Center",
                     "Request a quotation for European OEM and aftermarket parts. Send OE numbers, part numbers, quantities and VIN by email, or attach an Excel parts list.",
                     quote, "request-quote/"))

    # ---- Downloads ----
    def cat_cell(b):
        if b["catalog"]:
            return f'<a class="ext" href="{b["catalog"]}" target="_blank" rel="noopener">Open catalog ↗</a>'
        return '<a class="ext" href="{R}contact/">On request</a>'
    rows = "".join(
        f'<tr><td><a href="{{R}}brands/{b["slug"]}/"><b>{E(b["name"])}</b></a></td><td>{E(b["group"] or "—")}</td>'
        f'<td>{", ".join(SYSTEMS[s] for s in b["systems"])}</td><td>{cat_cell(b)}</td></tr>'
        for b in sorted(BRANDS, key=lambda b: b["name"].lower()))
    downloads = f'''
<section class="page-hero">
  <div class="container">
    <nav class="crumbs" aria-label="Breadcrumb"><a href="{{R}}">Home</a><span>/</span><span>Downloads</span></nav>
    <h1>Downloads</h1>
    <p>Our company profile, the parts-request template, and links to every brand's official catalog.</p>
  </div>
</section>
<section class="section">
  <div class="container">
    <div class="section-head"><span class="eyebrow">Company files</span><h2>Documents</h2></div>
    <div class="dl-list">
      <div class="dl reveal"><span class="dl__icon">{I["file"]}</span><div class="dl__body"><h3>Company Profile</h3><p>PDF · 100 KB · Our specialization, territory and contact details</p></div><a class="btn btn--primary" href="{{R}}{PDF}" download>{I["download"]}Download</a></div>
      <div class="dl reveal"><span class="dl__icon">{I["sheet"]}</span><div class="dl__body"><h3>Parts Request Template</h3><p>Excel · For bulk quotation requests (OE number, part number, brand, quantity, VIN)</p></div><a class="btn btn--primary" href="{{R}}{TEMPLATE_XLSX}" download>{I["download"]}Download</a></div>
      <div class="dl reveal"><span class="dl__icon">{I["db"]}</span><div class="dl__body"><h3>Data Sheet Samples</h3><p>Sample parts-data and cross-reference sheets are available to business customers on request.</p></div><a class="btn btn--outline" href="{{R}}contact/">Request samples</a></div>
    </div>
  </div>
</section>
<section class="section section--soft">
  <div class="container">
    <div class="section-head"><span class="eyebrow">Brand catalogs</span><h2>Official online catalogs</h2><p>Each link opens the manufacturer's own catalog, always up to date.</p></div>
    <div class="table-wrap"><table class="cat-table">
      <thead><tr><th>Brand</th><th>Group</th><th>Systems</th><th>Catalog</th></tr></thead>
      <tbody>{rows}</tbody>
    </table></div>
  </div>
</section>
'''
    urls.append(page("downloads/index.html", "Downloads | Euro Auto Center",
                     "Download the Euro Auto Center company profile and parts-request Excel template, and open official online catalogs for all authorized brands.",
                     downloads, "downloads/"))

    # ---- Contact ----
    locs = ""
    for o in OFFICES:
        q = quote_plus(o["map_q"])
        locs += f'''<article class="location reveal">
  <div class="location__info">
    <span class="tier">{o["label"]}</span>
    <h3>{o["city"]}</h3>
    <address>Euro Auto Center<br>{"<br>".join(E(l) for l in o["lines"])}</address>
    <div class="location__actions">
      <a class="btn btn--primary btn--sm" href="{o.get('maps_link') or ('https://www.google.com/maps/search/?api=1&amp;query=' + q)}" target="_blank" rel="noopener">{I["pin"]}Open in Google Maps</a>
      <a class="btn btn--outline btn--sm" href="https://www.google.com/maps/dir/?api=1&amp;destination={q}" target="_blank" rel="noopener">{I["route"]}Directions</a>
    </div>
  </div>
  <div class="map"><iframe title="Map: Euro Auto Center {o['city']}" src="{(o.get('embed') or ('https://www.google.com/maps?q=' + q + '&output=embed')).replace('&', '&amp;')}" loading="lazy" referrerpolicy="strict-origin-when-cross-origin" allowfullscreen></iframe></div>
</article>'''
    contact = f'''
<section class="page-hero">
  <div class="container">
    <nav class="crumbs" aria-label="Breadcrumb"><a href="{{R}}">Home</a><span>/</span><span>Contact</span></nav>
    <h1>Contact Us</h1>
    <p>Send us your part numbers, VIN or requirements. Our team replies with availability and pricing.</p>
  </div>
</section>
<section class="section">
  <div class="container">
    <div class="contact-grid">
      <a class="contact-card reveal" href="mailto:{SALES}"><span class="contact-card__icon">{I["mail"]}</span><span class="contact-card__label">Sales &amp; Quotes</span><span class="contact-card__value">{SALES}</span><span class="contact-card__hint">General and quotation requests</span></a>
      <a class="contact-card reveal" href="mailto:{DIRECT}"><span class="contact-card__icon">{I["user"]}</span><span class="contact-card__label">Commercial Contact</span><span class="contact-card__value">{DIRECT}</span><span class="contact-card__hint">Accounts and partnerships</span></a>
    </div>
  </div>
</section>
<section class="section section--soft">
  <div class="container">
    <div class="section-head"><span class="eyebrow">Our Location</span><h2>Visit our headquarters.</h2></div>
    {locs}
  </div>
</section>
<section class="section section--tight">
  <div class="container" style="text-align:center">
    <span class="eyebrow">Follow us</span>
    <h2 style="font-size:28px;margin:12px 0 18px">Connect with Euro Auto Center</h2>
    <div style="display:flex;justify-content:center;color:var(--navy-800)">{social_links("{R}")}</div>
  </div>
</section>
'''
    urls.append(page("contact/index.html", "Contact | Euro Auto Center",
                     f"Contact Euro Auto Center: {SALES}. Headquarters in Ras Al Khaimah, United Arab Emirates.",
                     contact, "contact/", ld_org()))

    # ---- assets ----
    os.makedirs(os.path.join(OUT, "assets/css")); os.makedirs(os.path.join(OUT, "assets/js")); os.makedirs(os.path.join(OUT, "assets/files"))
    shutil.copy(os.path.join(SRC, "site.css"), os.path.join(OUT, "assets/css/site.css"))
    shutil.copy(os.path.join(SRC, "site.js"), os.path.join(OUT, "assets/js/site.js"))
    for folder, dest in (("logos", "assets/brands"), ("social", "assets/social")):
        p = os.path.join(SRC, folder)
        if os.path.isdir(p) and os.listdir(p):
            shutil.copytree(p, os.path.join(OUT, dest))
    # The Excel template is edited by hand and lives in the repo; the build no longer regenerates it.

    # sitemap + robots
    with open(os.path.join(OUT, "sitemap.xml"), "w") as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?>\n<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">\n')
        for u in urls:
            f.write(f"  <url><loc>{u}</loc></url>\n")
        f.write("</urlset>\n")
    with open(os.path.join(OUT, "robots.txt"), "w") as f:
        f.write(f"User-agent: *\nAllow: /\n\nSitemap: {SITE}/sitemap.xml\n")
    print(f"Built {len(urls)} pages -> {OUT}")

def make_template(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "Parts Request"
    navy = PatternFill("solid", fgColor="0A1628")
    thin = Side(style="thin", color="D5DDE8")
    ws["A1"] = "EURO AUTO CENTER – Parts Quotation Request"
    ws["A1"].font = Font(bold=True, size=14, color="0A1628")
    ws["A2"] = f"Fill one part per row, save, and attach this file to your email to {SALES}"
    ws["A2"].font = Font(italic=True, size=10, color="4A5A72")
    headers = ["#", "OE Number", "Part Number", "Preferred Brand", "Description", "Quantity", "Vehicle (Make / Model / Year)", "VIN", "Notes"]
    widths = [5, 22, 20, 20, 32, 10, 30, 22, 30]
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = navy
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[4].height = 30
    example = [1, "1K0 615 301 AA", "", "TRW", "Brake disc, front", 2, "Volkswagen Golf 7, 2016", "", "Example row – replace or delete"]
    for i, v in enumerate(example, start=1):
        c = ws.cell(row=5, column=i, value=v)
        c.font = Font(italic=True, color="8A97AB")
    for r in range(5, 205):
        if r > 5:
            ws.cell(row=r, column=1, value=r - 4)
        for col in range(1, 10):
            ws.cell(row=r, column=col).border = Border(bottom=thin, left=thin, right=thin)
    dv = DataValidation(type="whole", operator="greaterThan", formula1="0", allow_blank=True,
                        errorTitle="Quantity", error="Please enter a whole number greater than 0.")
    ws.add_data_validation(dv); dv.add("F5:F204")
    blist = ",".join(["Any"] + sorted(b["name"] for b in BRANDS))
    if len(blist) < 255:
        dv2 = DataValidation(type="list", formula1=f'"{blist}"', allow_blank=True, showErrorMessage=False)
        ws.add_data_validation(dv2); dv2.add("D5:D204")
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "A4:I204"
    info = wb.create_sheet("Contact")
    for i, (k, v) in enumerate([("Email", SALES), ("Commercial", DIRECT), ("Headquarters", ", ".join(OFFICES[0]["lines"])),
                                ("Website", SITE)], start=1):
        info.cell(row=i, column=1, value=k).font = Font(bold=True)
        info.cell(row=i, column=2, value=v)
    info.column_dimensions["A"].width = 20; info.column_dimensions["B"].width = 60
    wb.save(path)

if __name__ == "__main__":
    build()
